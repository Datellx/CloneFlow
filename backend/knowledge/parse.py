"""Text extraction for uploaded documents and mail bodies (PDF, Word, PowerPoint, HTML, JSON, plain text)."""
from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

MAX_EVENT_CHARS = 12000


@dataclass
class RawEvent:
    kind: str
    locator: str
    text: str
    title: str = ""
    occurred_at: datetime | None = None
    image_paths: list[str] = field(default_factory=list)


def _clip(text: str) -> str:
    text = re.sub(r"\n{3,}", "\n\n", text.replace("\r\n", "\n")).strip()
    if len(text) > MAX_EVENT_CHARS:
        text = text[:MAX_EVENT_CHARS] + "\n[…truncated]"
    return text


def _norm_subject(s: str) -> str:
    return re.sub(r"^((re|fw|fwd|aw)\s*:\s*)+", "", (s or "").strip(), flags=re.I).lower()


def _msg_body(msg) -> str:
    if msg.is_multipart():
        for part in msg.walk():
            if part.get_content_type() == "text/plain" and not part.get_filename():
                try:
                    return part.get_content()
                except Exception:
                    return (part.get_payload(decode=True) or b"").decode("utf-8", "ignore")
        for part in msg.walk():
            if part.get_content_type() == "text/html":
                html = (part.get_payload(decode=True) or b"").decode("utf-8", "ignore")
                return re.sub(r"<[^>]+>", " ", html)
        return ""
    try:
        return msg.get_content()
    except Exception:
        return (msg.get_payload(decode=True) or b"").decode("utf-8", "ignore")


def _chunks(text: str, size: int = 6000) -> list[str]:
    paras = re.split(r"\n\s*\n", text)
    out, cur = [], ""
    for p in paras:
        if len(cur) + len(p) > size and cur:
            out.append(cur)
            cur = ""
        cur += p + "\n\n"
    if cur.strip():
        out.append(cur)
    return out or [text]


def parse_document(path: Path) -> list[RawEvent]:
    ext = path.suffix.lower()
    text = ""
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        text = "\n\n".join((p.extract_text() or "") for p in reader.pages)
    elif ext == ".docx":
        import docx

        d = docx.Document(str(path))
        text = "\n".join(p.text for p in d.paragraphs)
        for t in d.tables:
            for r in t.rows:
                text += "\n" + " | ".join(c.text for c in r.cells)
    elif ext == ".pptx":
        try:
            from pptx import Presentation

            prs = Presentation(str(path))
            text = "\n\n".join("\n".join(sh.text for sh in s.shapes if getattr(sh, "has_text_frame", False))
                               for s in prs.slides)
        except ImportError:
            text = ""
    elif ext in (".html", ".htm"):
        text = re.sub(r"<[^>]+>", " ", path.read_text(encoding="utf-8", errors="ignore"))
    elif ext == ".json":
        data = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
        text = json.dumps(data, indent=1)[:200000]
    else:
        text = path.read_text(encoding="utf-8", errors="ignore")
    if not text.strip():
        return []
    chunks = _chunks(text)
    return [RawEvent("document", f"{path.name}#part-{i + 1}" if len(chunks) > 1 else path.name, _clip(c), path.stem)
            for i, c in enumerate(chunks)]
